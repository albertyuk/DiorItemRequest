"""Extract highlighted SKUs from a Dior sell-thru map, look them up in a
stock query export, and fill a ProductsList template workbook.

All heavy workbooks are opened with ``read_only=True`` — the sell-thru map
expands to multiple GB of memory in normal mode.
"""

from __future__ import annotations

import base64
import io
import logging
import re
import shutil
from dataclasses import dataclass, field

from openpyxl import load_workbook

import sheet_images

log = logging.getLogger(__name__)

# --- Highlight recognition (kept configurable: a future recolor of the map
# should be a one-line change here) -----------------------------------------
# Yellow highlights in the sell-thru map are theme-indexed, not RGB:
# solid fill, fgColor.theme == 7 (accent4), tint ~= 0.80. fgColor.rgb is
# unset on those cells, so RGB-only detection finds nothing.
HIGHLIGHT_THEMES = {7}
# Fallback for workbooks that use literal RGB yellows instead of theme colors.
YELLOW_FAMILY = {"FFFF00", "FFF2CC", "FFE699", "FFC000", "FFD966"}
# Theme 4 is the structural light-blue banding on FW26 size rows — never a
# highlight and not worth flagging for review.
STRUCTURAL_THEMES = {4}

# SKU-shaped strings: 8+ alphanumerics, an X, then a 3+ alphanumeric suffix.
# Covers FW26 parent rows (X + 4-char color) and LOOK TOTAL variants.
SKU_RE = re.compile(r"^[0-9A-Z]{8,}X[0-9A-Z]{3,}$")
# The color block is always the *trailing* X + 4 alphanumerics (optionally
# separated, as in 641V19A1491_X8300). Never use split('X'): bases
# themselves may contain an X (652P92X3F74X8090).
COLOR_SUFFIX_RE = re.compile(r"[_\- ]?X[0-9A-Z]{4}$")

# Looser shape for values in AI-located SKU columns: codes the strict regex
# cannot recognize (accessory MMCs like M0759OWKAM912, underscore codes like
# KCK554TFS_S03W). Must contain both a letter and a digit.
LOOSE_SKU_RE = re.compile(r"^[0-9A-Za-z][0-9A-Za-z_\-./ ]{4,39}$")


def _plausible_sku(value: str) -> bool:
    if value.startswith("=") or not LOOSE_SKU_RE.match(value):
        return False
    return (any(ch.isdigit() for ch in value)
            and any(ch.isalpha() for ch in value)
            and value.count(" ") <= 1)

# --- Explicit "PickUp" selection ------------------------------------------
# A sheet may carry a marker column whose header normalizes to "pickup"
# (PickUp / pick up / PICK-UP ...). When present it is the PRIMARY selector
# for that sheet: rows marked with a recognized signal are extracted and
# highlights on that sheet are reported but no longer select. Sheets
# without the column keep the yellow-highlight behavior.
MARKER_HEADER_NORM = "pickup"
MARKER_HEADER_SCAN_ROWS = 15   # header must appear in the first N rows
# Canonical mark is "Y"; these case-insensitive equivalents also count.
MARKER_SIGNALS = {"y", "yes", "x", "1", "true", "ok", "pickup", "✓", "√",
                  "是"}


def _norm_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.casefold())


def _find_marker_column(ws) -> tuple[str | None, int]:
    """(column letter, header row) of the sheet's PickUp column, or
    (None, 0). Exact normalized match only — a stray word containing
    'pickup' must not hijack a sheet's selection."""
    for row in ws.iter_rows(min_row=1, max_row=MARKER_HEADER_SCAN_ROWS):
        for cell in row:
            if (isinstance(cell.value, str)
                    and _norm_header(cell.value) == MARKER_HEADER_NORM):
                return cell.column_letter, cell.row
    return None, 0


QUERY_SHEET = "query"
QUERY_HEADERS = [
    "Barcode", "Sku", "Division", "Department", "Season",
    "Product reference", "Color code", "Size", "Quantity",
    "Unit Weighted Average Cost", "Unit Retail Price incl. tax",
    "Item Type", "Path",
]

TEMPLATE_SHEET = "Sheet1"
TEMPLATE_COLS = 13  # A..M
IMAGE_COL = 14      # N — product images go all the way to the right
IMAGE_MAX_PX = 84   # thumbnail height in the output workbook

# Cap the "other fills" review list so one oddly formatted sheet cannot
# balloon the report.
MAX_OTHER_FILLS = 200
# Cap the per-sheet highlighted-cell trace shown on the report page (a fully
# highlighted sheet must not produce an unbounded page).
MAX_HIGHLIGHT_DETAILS = 1000


def is_highlighted(cell) -> bool:
    """True when the cell carries a yellow 'reorder' highlight."""
    f = cell.fill
    if f is None or f.patternType != "solid":
        return False
    fg = f.fgColor
    theme = getattr(fg, "theme", None)
    if isinstance(theme, int):
        return theme in HIGHLIGHT_THEMES
    rgb = fg.rgb if isinstance(fg.rgb, str) else ""
    return rgb[-6:].upper() in YELLOW_FAMILY


def _fill_note(cell) -> str | None:
    """Describe a solid fill that is neither a highlight nor structural
    banding, for the human review list. Returns None for anything else."""
    f = cell.fill
    if f is None or f.patternType != "solid":
        return None
    fg = f.fgColor
    theme = getattr(fg, "theme", None)
    if isinstance(theme, int):
        if theme in HIGHLIGHT_THEMES or theme in STRUCTURAL_THEMES:
            return None
        tint = fg.tint or 0
        return f"theme {theme}, tint {tint:.2f}"
    rgb = fg.rgb if isinstance(fg.rgb, str) else ""
    if rgb:
        if rgb[-6:].upper() in YELLOW_FAMILY:
            return None
        return f"rgb {rgb[-6:].upper()}"
    # Legacy color types (indexed palette, auto) — openpyxl exposes unset
    # attributes as descriptor objects, so nothing above matched. These are
    # never treated as highlights, but must still reach the human review
    # list rather than vanish silently.
    indexed = getattr(fg, "indexed", None)
    if isinstance(indexed, int):
        return f"indexed color {indexed}"
    return "unrecognized solid fill"


def extract_base(sku: str) -> str:
    """Strip the trailing color block (X + 4 alphanumerics) from a SKU."""
    return COLOR_SUFFIX_RE.sub("", sku.strip())


@dataclass
class SheetScan:
    name: str
    sku_cells: int = 0
    highlighted_cells: int = 0
    highlighted_skus: set[str] = field(default_factory=set)
    # step-by-step trace for the report page: (cell coordinate, sku, base)
    highlighted_cell_details: list[tuple[str, str, str]] = field(default_factory=list)
    other_fills: list[tuple[str, str, str]] = field(default_factory=list)  # (coord, sku, note)
    # cells found only through AI-located SKU columns (non-standard formats)
    ai_sku_cells: int = 0
    ai_highlighted_cell_details: list[tuple[str, str, str]] = field(default_factory=list)
    # explicit PickUp-column selection (primary when the column exists)
    marker_column: str | None = None
    marker_header_row: int = 0
    marked_rows: int = 0
    marked_skus: set[str] = field(default_factory=set)
    marked_cell_details: list[tuple[str, str, str]] = field(default_factory=list)
    marker_unrecognized: list[tuple[str, str]] = field(default_factory=list)  # (coord, value)
    marker_no_sku_rows: list[str] = field(default_factory=list)  # marker coords

    @property
    def uses_marker(self) -> bool:
        return self.marker_column is not None

    @property
    def selected_skus(self) -> set[str]:
        """The sheet's selection: PickUp marks when the column exists,
        else the yellow highlights."""
        return self.marked_skus if self.uses_marker else self.highlighted_skus

    @property
    def highlighted_unmarked(self) -> list[str]:
        """Highlighted SKUs a PickUp sheet did NOT mark — shown to the
        human, because highlights alone no longer select there."""
        if not self.uses_marker:
            return []
        return sorted(self.highlighted_skus - self.marked_skus)

    @property
    def bases(self) -> set[str]:
        return {extract_base(s) for s in self.selected_skus}


@dataclass
class MapScan:
    sheets: list[SheetScan]

    @property
    def bases(self) -> dict[str, list[str]]:
        """Union of highlighted bases -> sorted list of sheets they appear on."""
        out: dict[str, set[str]] = {}
        for sheet in self.sheets:
            for base in sheet.bases:
                out.setdefault(base, set()).add(sheet.name)
        return {b: sorted(names) for b, names in sorted(out.items())}

    @property
    def other_fills(self) -> list[tuple[str, str, str, str]]:
        return [
            (sheet.name, coord, sku, note)
            for sheet in self.sheets
            for coord, sku, note in sheet.other_fills
        ]


def scan_sell_thru_map(path, ai_columns: dict | None = None,
                       progress=None) -> MapScan:
    """Scan every cell of every sheet for SKU-shaped strings and classify
    their fills. Geometry-free on purpose: it handles both the FW26 row
    layout and the LOOK TOTAL block layout without hardcoded columns.

    ai_columns (optional) extends the scan with AI-located SKU columns:
    {sheet name: {"header_row": int | None, "columns": {letter: header}}}.
    In those columns, plausible SKU codes that the strict regex cannot
    recognize are also extracted (below the header row only).

    progress (optional): callback(stage, done, total) invoked periodically.
    """
    ai_columns = ai_columns or {}
    wb = load_workbook(path, read_only=True)
    try:
        total_rows = sum((ws.max_row or 0) for ws in wb.worksheets) or 1
        seen_rows = 0
        sheets = []
        for ws in wb.worksheets:
            scan = SheetScan(name=ws.title)
            scan.marker_column, scan.marker_header_row = \
                _find_marker_column(ws)
            m_col, m_row = scan.marker_column, scan.marker_header_row
            sheet_ai = ai_columns.get(ws.title) or {}
            ai_cols = sheet_ai.get("columns") or {}
            ai_min_row = (sheet_ai.get("header_row") or 0) + 1
            for row in ws.iter_rows():
                seen_rows += 1
                if progress and seen_rows % 200 == 0:
                    progress("scan", seen_rows, total_rows)
                # Marker state first: it decides how this row's SKUs count.
                marked = False
                marker_coord = None
                if m_col:
                    for cell in row:
                        if cell.value is None:
                            continue
                        if (cell.column_letter == m_col
                                and cell.row > m_row):
                            raw = str(cell.value).strip()
                            marker_coord = cell.coordinate
                            if raw.casefold() in MARKER_SIGNALS:
                                marked = True
                                scan.marked_rows += 1
                            elif len(scan.marker_unrecognized) < MAX_OTHER_FILLS:
                                scan.marker_unrecognized.append(
                                    (cell.coordinate, raw[:40]))
                            break
                row_marked_sku = False
                for cell in row:
                    value = cell.value
                    if not isinstance(value, str):
                        continue
                    if m_col and cell.column_letter == m_col:
                        continue  # the marker column never holds SKUs
                    sku = value.strip()
                    strict = bool(SKU_RE.match(sku))
                    in_ai = (not strict and bool(ai_cols)
                             and cell.row >= ai_min_row
                             and cell.column_letter in ai_cols
                             and _plausible_sku(sku))
                    if not strict and not in_ai:
                        continue
                    scan.sku_cells += 1
                    if in_ai:
                        scan.ai_sku_cells += 1
                    highlighted = is_highlighted(cell)
                    if highlighted:
                        scan.highlighted_cells += 1
                        scan.highlighted_skus.add(sku)
                        details = (scan.highlighted_cell_details if strict
                                   else scan.ai_highlighted_cell_details)
                        if len(details) < MAX_HIGHLIGHT_DETAILS:
                            details.append(
                                (cell.coordinate, sku, extract_base(sku)))
                    elif strict and len(scan.other_fills) < MAX_OTHER_FILLS:
                        note = _fill_note(cell)
                        if note:
                            scan.other_fills.append(
                                (cell.coordinate, sku, note))
                    if marked:
                        scan.marked_skus.add(sku)
                        if len(scan.marked_cell_details) < MAX_HIGHLIGHT_DETAILS:
                            scan.marked_cell_details.append(
                                (cell.coordinate, sku, extract_base(sku)))
                        if in_ai and not highlighted:
                            # AI-located SKU selected by mark: record it in
                            # the AI trace too (column-memory approval and
                            # the report's AI section both read this list).
                            if len(scan.ai_highlighted_cell_details) < MAX_HIGHLIGHT_DETAILS:
                                scan.ai_highlighted_cell_details.append(
                                    (cell.coordinate, sku, extract_base(sku)))
                        row_marked_sku = True
                if (marked and not row_marked_sku and marker_coord
                        and len(scan.marker_no_sku_rows) < MAX_OTHER_FILLS):
                    scan.marker_no_sku_rows.append(marker_coord)
            sheets.append(scan)
        if progress:
            progress("scan", total_rows, total_rows)
        return MapScan(sheets=sheets)
    finally:
        wb.close()


def read_query_headers(path) -> list:
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb[QUERY_SHEET] if QUERY_SHEET in wb.sheetnames else wb.worksheets[0]
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = list(first)
        while headers and headers[-1] is None:
            headers.pop()
        return headers
    finally:
        wb.close()


def validate_query_file(path) -> tuple[bool, str, int]:
    """Check the 13 expected headers; returns (ok, message, data_row_count)."""
    try:
        headers = read_query_headers(path)
    except Exception as exc:  # malformed zip/xlsx
        return False, f"could not be read as an .xlsx workbook ({exc})", 0
    normalized = [str(h).strip() if h is not None else "" for h in headers]
    if normalized != QUERY_HEADERS:
        return (
            False,
            "header row does not match the expected 13 query columns "
            f"(got: {normalized[:14]})",
            0,
        )
    return True, "ok", count_query_rows(path)


def count_query_rows(path) -> int:
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb[QUERY_SHEET] if QUERY_SHEET in wb.sheetnames else wb.worksheets[0]
        return sum(
            1
            for row in ws.iter_rows(min_row=2, values_only=True)
            if any(v is not None for v in row)
        )
    finally:
        wb.close()


def match_query(path, bases, progress=None) -> dict[str, list[tuple]]:
    """Single pass over the query export. A query Sku is 'BASE - COLOR - SIZE'
    with literal ' - ' separators; match on exact first-segment equality
    (never prefix-match raw strings — first segments can contain X too).

    Returns {base: [query row tuples]} containing only matched bases.
    """
    bases = set(bases)
    matched: dict[str, list[tuple]] = {}
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb[QUERY_SHEET] if QUERY_SHEET in wb.sheetnames else wb.worksheets[0]
        total = ws.max_row or 1
        for row_num, row in enumerate(
                ws.iter_rows(min_row=2, values_only=True), start=2):
            if progress and row_num % 2000 == 0:
                progress("match", row_num, total)
            sku = row[1] if len(row) > 1 else None
            if not isinstance(sku, str):
                continue
            base = sku.split(" - ")[0]
            if base in bases:
                if len(row) < 13:  # ragged row: pad so consumers can unpack
                    row = tuple(row) + (None,) * (13 - len(row))
                matched.setdefault(base, []).append(row)
    finally:
        wb.close()
    return matched


def _sku_sort_key(row) -> tuple[str, str, str]:
    parts = [p.strip() for p in str(row[1]).split(" - ")]
    parts += [""] * (3 - len(parts))
    return parts[0], parts[1], parts[2]


def _barcode_text(value) -> str:
    """Barcodes must stay text — no lost leading zeros, no scientific
    notation. The query stores them as strings already; be defensive about
    numeric cells anyway."""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if value is None:
        return ""
    return str(value)


def _place_image(ws, row: int, data: bytes) -> bool:
    """Anchor one product thumbnail in the image column of `row`."""
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.utils import get_column_letter

    try:
        img = XLImage(io.BytesIO(data))
    except Exception:  # not a raster PIL can open — skip, never fail a run
        return False
    if not img.height or not img.width:
        return False
    scale = min(1.0, IMAGE_MAX_PX / img.height)
    img.height = int(img.height * scale)
    img.width = int(img.width * scale)
    ws.add_image(img, f"{get_column_letter(IMAGE_COL)}{row}")
    # px -> pt (0.75) plus a little breathing room
    ws.row_dimensions[row].height = max(
        ws.row_dimensions[row].height or 15, int(img.height * 0.75) + 6)
    return True


def build_output(template_path, out_path, matched, unmatched_bases,
                 base_sheets=None, images=None, progress=None) -> int:
    """Copy the template and append one row per matched query row, sorted by
    base, then color, then size. Returns the number of data rows written.

    The query and template columns do NOT align — rows are mapped by meaning,
    and query's Product reference / Color code / Size / Item Type / Path are
    intentionally dropped.

    images (optional): {base: raw image bytes}. Each base's product picture
    is placed once, on the first row of its block, in column N — all the
    way to the right of the template's 13 columns.
    """
    shutil.copyfile(template_path, out_path)
    wb = load_workbook(out_path)
    ws = wb[TEMPLATE_SHEET] if TEMPLATE_SHEET in wb.sheetnames else wb.worksheets[0]

    # Row 2 of the template only demonstrates the H/J formula pattern; clear
    # it so it is either overwritten as the first data row or not left behind.
    for col in range(1, TEMPLATE_COLS + 1):
        ws.cell(row=2, column=col).value = None

    rows = [row for base in sorted(matched) for row in matched[base]]
    rows.sort(key=_sku_sort_key)

    images = images or {}
    if images:
        ws.cell(row=1, column=IMAGE_COL, value="Image")
    placed: set[str] = set()

    total = len(rows) or 1
    n = 1
    for query_row in rows:
        n += 1
        if progress and (n - 2) % 100 == 0:
            progress("write", n - 2, total)
        (barcode, sku, division, department, season,
         _ref, _color, _size, quantity, unit_cost, unit_retail,
         _item_type, _path) = query_row[:13]
        barcode_cell = ws.cell(row=n, column=1, value=_barcode_text(barcode))
        barcode_cell.number_format = "@"
        ws.cell(row=n, column=2, value=sku)
        ws.cell(row=n, column=3, value=division)
        ws.cell(row=n, column=4, value=department)
        ws.cell(row=n, column=5, value=season)
        ws.cell(row=n, column=6, value=quantity)
        ws.cell(row=n, column=7, value=unit_cost)
        ws.cell(row=n, column=8,
                value=f'=IF(AND(F{n}<>"", G{n}<>""), F{n}*G{n}, "")')
        ws.cell(row=n, column=9, value=unit_retail)
        ws.cell(row=n, column=10,
                value=f'=IF(AND(F{n}<>"", I{n}<>""), F{n}*I{n}, "")')
        # K, L, M (Gift Recipient / Organization / Position) stay blank.
        base = str(sku).split(" - ")[0]
        if base in images and base not in placed:
            if _place_image(ws, n, images[base]):
                placed.add(base)

    unmatched_ws = wb.create_sheet("Unmatched")
    unmatched_ws.append(["Base SKU", "Highlighted on sheet(s)"])
    base_sheets = base_sheets or {}
    for base in sorted(unmatched_bases):
        unmatched_ws.append([base, ", ".join(base_sheets.get(base, []))])

    wb.save(out_path)
    wb.close()
    return n - 1


@dataclass
class RunReport:
    sheets: list                         # SheetScan objects or plain dicts
    bases: dict[str, list[str]]          # base -> sheets it was highlighted on
    base_skus: dict[str, list[str]]      # base -> highlighted colorway SKUs
    matched_counts: dict[str, int]       # base -> number of query rows
    matched_rows: dict[str, list[dict]]  # base -> query data pulled (trace)
    unmatched: list[str]
    other_fills: list
    rows_written: int
    ai_enabled: bool = False             # was an AI column locator supplied?
    ai_detection: list | None = None     # locator output per sheet
    ai_note: str | None = None           # why detection is missing/failed
    ai_confirmed: bool = False           # did a human review the run?
    excluded: list | None = None         # bases the human unticked at review
    base_sources: dict | None = None     # base -> ["standard"|"ai", ...]
    query_info: dict | None = None       # stored-query metadata at scan time
    images_count: int = 0                # product images placed in column N


def sheet_to_dict(s: SheetScan) -> dict:
    """JSON-serializable form of a sheet scan (drafts, report files)."""
    return {
        "name": s.name,
        "sku_cells": s.sku_cells,
        "highlighted_cells": s.highlighted_cells,
        "highlighted_skus": sorted(s.highlighted_skus),
        "bases": sorted(s.bases),
        "highlighted_cell_details": s.highlighted_cell_details,
        "other_fills": s.other_fills,
        "ai_sku_cells": s.ai_sku_cells,
        "ai_highlighted_cell_details": s.ai_highlighted_cell_details,
        "marker_column": s.marker_column,
        "marker_header_row": s.marker_header_row,
        "marked_rows": s.marked_rows,
        "marked_cell_details": s.marked_cell_details,
        "marker_unrecognized": s.marker_unrecognized,
        "marker_no_sku_rows": s.marker_no_sku_rows,
        "highlighted_unmarked": s.highlighted_unmarked[:200],
    }


def build_ai_columns(detection) -> dict:
    """Reshape locator output for scan_sell_thru_map's ai_columns input."""
    out = {}
    for entry in detection or []:
        columns = {c["column"]: c.get("header") for c in entry["sku_columns"]}
        if columns:
            out[entry["sheet"]] = {
                "header_row": entry.get("header_row"),
                "columns": columns,
            }
    return out


def _trace_rows(rows) -> list[dict]:
    """Shape matched query rows for the report page's step-by-step trace."""
    out = []
    for r in sorted(rows, key=_sku_sort_key):
        out.append({
            "barcode": _barcode_text(r[0]), "sku": r[1], "division": r[2],
            "department": r[3], "season": r[4], "quantity": r[8],
            "unit_cost": r[9], "unit_retail": r[10],
        })
    return out


def scan_and_match(map_path, query_path, ai_detection=None,
                   progress=None) -> dict:
    """First half of a run: extract highlighted SKUs (standard scan plus any
    AI-located columns) and pull their query rows. Returns a
    JSON-serializable draft that the human review step can filter before the
    output is built — matching happens HERE, so the draft is immune to the
    stored query being replaced while the review page sits open."""
    scan = scan_sell_thru_map(map_path,
                              ai_columns=build_ai_columns(ai_detection),
                              progress=progress)
    bases = scan.bases
    matched = match_query(query_path, bases, progress=progress)

    base_skus: dict[str, set] = {}
    sources: dict[str, set] = {}
    for sheet in scan.sheets:
        for sku in sheet.selected_skus:
            base_skus.setdefault(extract_base(sku), set()).add(sku)
        if sheet.uses_marker:
            for _, _, base in sheet.marked_cell_details:
                sources.setdefault(base, set()).add("pickup")
            for _, _, base in sheet.ai_highlighted_cell_details:
                sources.setdefault(base, set()).add("ai")
        else:
            for _, _, base in sheet.highlighted_cell_details:
                sources.setdefault(base, set()).add("standard")
            for _, _, base in sheet.ai_highlighted_cell_details:
                sources.setdefault(base, set()).add("ai")

    # Product images: pull them for the selected SKU rows NOW — the map
    # file is deleted right after the scan, so the draft must carry them
    # (base64) through the review checkpoint to the build.
    wanted: dict[str, dict[int, str]] = {}
    for sheet in scan.sheets:
        details = (sheet.marked_cell_details if sheet.uses_marker
                   else (sheet.highlighted_cell_details
                         + sheet.ai_highlighted_cell_details))
        rows_map = sheet_images.rows_by_base(details)
        if rows_map:
            wanted[sheet.name] = rows_map

    return {
        "sheets": [sheet_to_dict(s) for s in scan.sheets],
        "bases": bases,
        "base_skus": {b: sorted(v) for b, v in sorted(base_skus.items())},
        "base_sources": {b: sorted(v) for b, v in sorted(sources.items())},
        "matched": {b: [list(r) for r in rows]
                    for b, rows in matched.items()},
        "other_fills": scan.other_fills,
        "images": sheet_images.extract_for(map_path, wanted),
    }


def build_from_selection(template_path, out_path, draft, selected=None,
                         ai_info: dict | None = None,
                         progress=None) -> RunReport:
    """Second half of a run: build the ProductsList workbook from a draft,
    keeping only the bases the human left selected (None = keep all)."""
    bases = draft["bases"]
    if selected is None:
        kept = set(bases)
    else:
        kept = set(selected) & set(bases)
    excluded = sorted(set(bases) - kept)

    matched = {b: [tuple(r) for r in rows]
               for b, rows in draft["matched"].items() if b in kept}
    unmatched = sorted(b for b in kept if b not in matched)

    # Decode + validate the product images for the kept, matched bases
    # (only decodable rasters count, so the report number is exact).
    images: dict[str, bytes] = {}
    for b in set(matched) & set(draft.get("images") or {}):
        try:
            from PIL import Image as PILImage
            data = base64.b64decode(draft["images"][b]["data"])
            PILImage.open(io.BytesIO(data)).verify()
            images[b] = data
        except Exception:
            continue

    rows_written = build_output(
        template_path, out_path, matched, unmatched, base_sheets=bases,
        images=images, progress=progress)

    ai_info = ai_info or {}
    report = RunReport(
        sheets=draft["sheets"],
        bases=bases,
        base_skus=draft.get("base_skus", {}),
        matched_counts={b: len(matched[b]) for b in sorted(matched)},
        matched_rows={b: _trace_rows(matched[b]) for b in sorted(matched)},
        unmatched=unmatched,
        other_fills=[tuple(f) for f in draft.get("other_fills", [])],
        rows_written=rows_written,
        ai_enabled=bool(ai_info.get("enabled")),
        ai_detection=ai_info.get("detection"),
        ai_note=ai_info.get("note"),
        ai_confirmed=bool(ai_info.get("confirmed")),
        excluded=excluded,
        base_sources=draft.get("base_sources"),
        query_info=draft.get("query_info"),
        images_count=len(images),
    )
    log.info(
        "run: bases=%d kept=%d matched=%d unmatched=%d excluded=%d rows=%d",
        len(bases), len(kept), len(report.matched_counts),
        len(report.unmatched), len(excluded), report.rows_written,
    )
    return report


def run_pipeline(map_path, query_path, template_path, out_path,
                 locator=None, ai_result: dict | None = None,
                 progress=None) -> RunReport:
    """One-shot run with every extracted SKU kept (no human filtering).

    locator (optional): callable(map_path) -> detection list, as returned
    by sku_locator.locate. Any locator failure is reported, never fatal.
    ai_result (optional) supplies pre-computed detection state instead.
    """
    if ai_result is not None:
        ai_info = {
            "enabled": bool(ai_result.get("enabled")),
            "detection": ai_result.get("detection"),
            "note": ai_result.get("note"),
            "confirmed": bool(ai_result.get("confirmed")),
        }
    else:
        ai_info = {"enabled": locator is not None, "detection": None,
                   "note": None, "confirmed": False}
        if locator is not None:
            try:
                ai_info["detection"] = locator(map_path)
            except Exception as exc:
                log.warning("SKU column detection unavailable: %s", exc)
                ai_info["note"] = str(exc)
    draft = scan_and_match(map_path, query_path,
                           ai_detection=ai_info.get("detection"),
                           progress=progress)
    return build_from_selection(template_path, out_path, draft,
                                selected=None, ai_info=ai_info,
                                progress=progress)
