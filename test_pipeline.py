"""Acceptance tests against the real sample workbooks in ./samples/ plus
generated fixtures for layouts the sample copy does not contain.

The real samples hold internal pricing data and are not committed; tests that
need them skip when ./samples/ is absent.
"""

from __future__ import annotations

import io
import re
import time
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Color, PatternFill

import column_memory
import pipeline
import sku_locator
from app import create_app

SAMPLES = Path(__file__).parent / "samples"
MAP = SAMPLES / "sell_thru_map.xlsx"
QUERY = SAMPLES / "query.xlsx"
TEMPLATE = SAMPLES / "ProductsListTemplate.xlsx"

needs_samples = pytest.mark.skipif(
    not (MAP.exists() and QUERY.exists() and TEMPLATE.exists()),
    reason="real sample workbooks not present in ./samples/",
)


def _wait_status(client, job_id, timeout=180):
    """Poll the progress endpoint until the background job finishes."""
    statuses = []
    deadline = time.time() + timeout
    while time.time() < deadline:
        data = client.get(f"/progress/{job_id}/status").get_json()
        statuses.append(data)
        if data["status"] in ("done", "error"):
            return data, statuses
        time.sleep(0.2)
    raise AssertionError(f"job {job_id} did not finish: {statuses[-3:]}")


def _finish_job(client, resp, timeout=180):
    """Follow a 302 to /progress/<job_id> and wait out the job."""
    assert resp.status_code == 302, (resp.status_code, resp.data[:300])
    location = resp.headers["Location"]
    match = re.search(r"/progress/([0-9a-f_]+)$", location)
    assert match, location
    return _wait_status(client, match.group(1), timeout)


@pytest.fixture(scope="session")
def map_scan():
    return pipeline.scan_sell_thru_map(MAP)


@pytest.fixture(scope="session")
def fw26_matched(map_scan):
    fw26 = next(s for s in map_scan.sheets if s.name == "FW26")
    return pipeline.match_query(QUERY, fw26.bases)


# --- 1. highlight extraction counts ----------------------------------------

@needs_samples
def test_fw26_highlight_counts(map_scan):
    fw26 = next(s for s in map_scan.sheets if s.name == "FW26")
    assert fw26.highlighted_cells == 43
    assert len(fw26.bases) == 41
    # the structural theme-4 banding on size rows must never be flagged
    assert fw26.other_fills == []


@needs_samples
def test_look_total_highlight_counts(map_scan):
    lt = next((s for s in map_scan.sheets if s.name == "LOOK TOTAL"), None)
    if lt is None:
        pytest.skip(
            "this sample copy of the map has no 'LOOK TOTAL' sheet "
            "(only FW26 and an empty '-->' sheet); the block layout is "
            "covered by test_look_total_block_layout_synthetic"
        )
    assert lt.highlighted_cells == 60
    assert len(lt.bases) == 39


def _yellow():
    return PatternFill(fill_type="solid", start_color=Color(theme=7, tint=0.8))


def test_look_total_block_layout_synthetic(tmp_path):
    """Replicates the documented LOOK TOTAL geometry: 8-column look blocks
    starting at A, J, S, AB, ... with SKUs under 'MMC' (block-relative col 3),
    data from row 8; 60 highlighted cells over 46 unique SKUs / 39 bases,
    plus gray-filled cells that must be flagged, not extracted."""
    bases = [f"644T{i:03d}A99" for i in range(39)]
    skus = [b + "X5800" for b in bases] + [b + "X0900" for b in bases[:7]]
    assert len(set(skus)) == 46
    highlight_cells = skus + skus[:14]  # 60 highlighted cells
    plain_cells = [b + "X5800" for b in bases[:20]]

    wb = Workbook()
    ws = wb.active
    ws.title = "LOOK TOTAL"
    block_cols = [3, 12, 21, 30, 39, 48, 57, 66]  # C, L, U, AD, AM, AV, BE, BN
    for col in block_cols:
        ws.cell(row=7, column=col, value="MMC")

    slots = [(row, col) for row in range(8, 40) for col in block_cols]
    it = iter(slots)
    for sku in highlight_cells:
        row, col = next(it)
        c = ws.cell(row=row, column=col, value=sku)
        c.fill = _yellow()
    for sku in plain_cells:
        row, col = next(it)
        ws.cell(row=row, column=col, value=sku)
    gray_specs = [(0, -0.15), (0, -0.15), (0, -0.15), (6, 0.8), (6, 0.8)]
    for i, (theme, tint) in enumerate(gray_specs):
        row, col = next(it)
        c = ws.cell(row=row, column=col, value=bases[20 + i] + "X5800")
        c.fill = PatternFill(fill_type="solid",
                             start_color=Color(theme=theme, tint=tint))
    # structural banding (theme 4) must be silently ignored
    row, col = next(it)
    c = ws.cell(row=row, column=col, value=bases[30] + "X5800")
    c.fill = PatternFill(fill_type="solid",
                         start_color=Color(theme=4, tint=0.8))

    path = tmp_path / "look_total.xlsx"
    wb.save(path)

    scan = pipeline.scan_sell_thru_map(path)
    lt = next(s for s in scan.sheets if s.name == "LOOK TOTAL")
    assert lt.highlighted_cells == 60
    assert len(lt.highlighted_skus) == 46
    assert len(lt.bases) == 39
    grays = [note for _, _, note in lt.other_fills]
    assert len(grays) == 5
    assert all(note.startswith("theme") for note in grays)
    # the step-trace records where each highlight was found and its base
    assert len(lt.highlighted_cell_details) == 60
    assert lt.highlighted_cell_details[0] == ("C8", "644T000A99X5800", "644T000A99")


def test_rgb_yellow_fallback(tmp_path):
    """A future map that uses literal RGB yellows must still be detected."""
    wb = Workbook()
    ws = wb.active
    c = ws.cell(row=1, column=1, value="644T000A99X5800")
    c.fill = PatternFill(fill_type="solid", start_color="FFFFF2CC")
    path = tmp_path / "rgb.xlsx"
    wb.save(path)
    scan = pipeline.scan_sell_thru_map(path)
    assert scan.sheets[0].highlighted_cells == 1


# --- 2. base-SKU extraction -------------------------------------------------

def test_base_extraction_double_x():
    assert pipeline.extract_base("652P92X3F74X8090") == "652P92X3F74"
    assert re.sub(r"X[0-9A-Z]{4}$", "", "652P92X3F74X8090") == "652P92X3F74"
    assert pipeline.extract_base("644S83A7A26X5883") == "644S83A7A26"
    assert pipeline.extract_base("  644S83A7A26X5883 ") == "644S83A7A26"
    # underscore-separated color codes (delivery-tracker style)
    assert pipeline.extract_base("641V19A1491_X8300") == "641V19A1491"
    # no recognizable color block -> base is the SKU itself
    assert pipeline.extract_base("KCK554TFS_S03W") == "KCK554TFS_S03W"


# --- 2b. AI-located SKU columns ---------------------------------------------

def _track_workbook(path):
    """Delivery-tracker-style sheet: SKUs under a 'SKU' header in a format
    the strict regex cannot recognize."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Track"
    ws.append(["SKU", "QTY"])
    c = ws.cell(row=2, column=1, value="M0715OUQO_M900_TU")
    c.fill = _yellow()
    ws.cell(row=3, column=1, value="KCV536VCR_S900_T37")
    c = ws.cell(row=4, column=1, value="1222")  # numeric — not a SKU
    c.fill = _yellow()
    ws.cell(row=5, column=1, value="641V19A1491_X8300").fill = _yellow()
    wb.save(path)


def test_ai_columns_extend_scan(tmp_path):
    path = tmp_path / "track.xlsx"
    _track_workbook(path)

    # without AI columns the strict regex finds nothing here
    plain = pipeline.scan_sell_thru_map(path).sheets[0]
    assert plain.sku_cells == 0 and plain.highlighted_cells == 0

    ai_columns = {"Track": {"header_row": 1, "columns": {"A": "SKU"}}}
    scan = pipeline.scan_sell_thru_map(path, ai_columns=ai_columns).sheets[0]
    assert scan.ai_sku_cells == 3          # '1222' filtered as implausible
    assert scan.highlighted_cells == 2     # M0715... and 641V19...
    assert ("A2", "M0715OUQO_M900_TU", "M0715OUQO_M900_TU") \
        in scan.ai_highlighted_cell_details
    # underscore color code is stripped for the base
    assert ("A5", "641V19A1491_X8300", "641V19A1491") \
        in scan.ai_highlighted_cell_details
    # the header cell itself is never extracted
    assert all(coord != "A1" for coord, _, _ in scan.ai_highlighted_cell_details)


def test_pipeline_with_fake_locator(tmp_path):
    map_path = tmp_path / "map.xlsx"
    _track_workbook(map_path)
    qwb = Workbook()
    qws = qwb.active
    qws.title = "query"
    qws.append(pipeline.QUERY_HEADERS)
    qws.append(["3617000000001", "641V19A1491 - X8300 - T36", "d", "dep",
                "s", None, None, None, 1, 10, 20, "t", "p"])
    query_path = tmp_path / "q.xlsx"
    qwb.save(query_path)

    detection = [{"sheet": "Track", "header_row": 1,
                  "sku_columns": [{"column": "A", "header": "SKU",
                                   "reason": "data cells hold item codes"}]}]
    report = pipeline.run_pipeline(
        map_path, query_path,
        Path(__file__).parent / "assets" / "ProductsListTemplate.xlsx",
        tmp_path / "out.xlsx",
        locator=lambda p: detection,
    )
    assert report.ai_enabled and report.ai_detection == detection
    assert report.ai_note is None
    # 641V19A1491 matched the query; the underscore code did not
    assert report.matched_counts == {"641V19A1491": 1}
    assert "M0715OUQO_M900_TU" in report.unmatched
    assert report.rows_written == 1


def test_pipeline_survives_locator_failure(tmp_path):
    map_path = tmp_path / "map.xlsx"
    _track_workbook(map_path)
    qwb = Workbook()
    qws = qwb.active
    qws.title = "query"
    qws.append(pipeline.QUERY_HEADERS)
    query_path = tmp_path / "q.xlsx"
    qwb.save(query_path)

    def broken(_path):
        raise RuntimeError("api unreachable")

    report = pipeline.run_pipeline(
        map_path, query_path,
        Path(__file__).parent / "assets" / "ProductsListTemplate.xlsx",
        tmp_path / "out.xlsx",
        locator=broken,
    )
    assert report.ai_enabled
    assert report.ai_note == "api unreachable"
    assert report.rows_written == 0  # run completed regardless


def test_locator_previews_and_normalization(tmp_path):
    path = tmp_path / "track.xlsx"
    _track_workbook(path)
    previews = sku_locator.build_previews(path)
    assert previews[0]["sheet"] == "Track"
    assert previews[0]["rows"]["1"]["A"] == "SKU"
    assert previews[0]["rows"]["2"]["A"] == "M0715OUQO_M900_TU"

    # rows whose first cell is empty must not crash: read-only empty cells
    # are a singleton with no coordinates (regression: 'EmptyCell' object
    # has no attribute 'row')
    wb = Workbook()
    ws = wb.active
    ws.title = "Gap"
    ws["B1"] = "SKU"            # column A entirely empty
    ws["B3"] = "M0715OUQO_M900_TU"   # row 2 entirely empty
    gap = tmp_path / "gap.xlsx"
    wb.save(gap)
    previews = sku_locator.build_previews(gap)
    assert previews == [{"sheet": "Gap",
                         "rows": {"1": {"B": "SKU"},
                                  "3": {"B": "M0715OUQO_M900_TU"}}}]

    raw = [
        {"sheet": "Track", "header_row": 1,
         "sku_columns": [{"column": "a", "header": "SKU", "reason": "ok"},
                         {"column": "5", "header": None, "reason": "bad"}]},
        {"sheet": "Ghost", "header_row": 1, "sku_columns": []},
    ]
    normalized = sku_locator._normalize(raw, ["Track"])
    assert normalized == [{"sheet": "Track", "header_row": 1,
                           "sku_columns": [{"column": "A", "header": "SKU",
                                            "reason": "ok"}],
                           "source": "ai"}]
    # sheets where the model found no SKU columns are dropped entirely —
    # the realistic "nothing found" answer is per-sheet empty lists, and
    # it must not read as a positive detection
    assert sku_locator._normalize(
        [{"sheet": "Track", "header_row": 2, "sku_columns": []}],
        ["Track"]) == []


def _verification_app(tmp_path, monkeypatch, detection):
    """App with a faked AI locator: configured, returns `detection`."""
    monkeypatch.setattr(sku_locator, "is_configured", lambda: True)
    monkeypatch.setattr(sku_locator, "locate_from_previews",
                        lambda previews: detection)
    app = create_app(data_dir=tmp_path / "data", open_access=True)
    map_path = tmp_path / "track.xlsx"
    _track_workbook(map_path)
    qwb = Workbook()
    qws = qwb.active
    qws.title = "query"
    qws.append(pipeline.QUERY_HEADERS)
    qws.append(["3617000000001", "641V19A1491 - X8300 - T36", "d", "dep",
                "s", None, None, None, 1, 10, 20, "t", "p"])
    query_path = tmp_path / "q.xlsx"
    qwb.save(query_path)
    return app.test_client(), map_path, query_path


def _upload(client, map_path, query_path, extra_query=True):
    if extra_query:
        with map_path.open("rb") as m, query_path.open("rb") as q:
            return client.post(
                "/process",
                data={"map_file": (m, "track.xlsx"),
                      "query_file": (q, "q.xlsx")},
                content_type="multipart/form-data",
            )
    with map_path.open("rb") as m:
        return client.post("/process", data={"map_file": (m, "track.xlsx")},
                           content_type="multipart/form-data")


def test_review_flow_with_ai_columns(tmp_path, monkeypatch):
    """Scan lands on the review checkpoint; the user keeps everything."""
    detection = [{"sheet": "Track", "header_row": 1,
                  "sku_columns": [{"column": "A", "header": "SKU",
                                   "reason": "data cells hold item codes"}]}]
    client, map_path, query_path = _verification_app(
        tmp_path, monkeypatch, detection)

    data, _ = _finish_job(client, _upload(client, map_path, query_path))
    assert data["status"] == "done", data
    assert "/review/" in data["next"]

    html = client.get(data["next"]).data.decode()
    assert "Review before building" in html
    assert "M0715OUQO_M900_TU" in html          # AI-found SKU listed
    assert "badge-ai" in html                    # marked as AI-sourced
    assert "no stock" in html                    # unmatched badge shown
    skus = re.findall(r'name="sku" value="([^"]+)"', html)
    # only HIGHLIGHTED cells become bases (KCV536VCR_S900_T37 has no fill)
    assert set(skus) == {"M0715OUQO_M900_TU", "641V19A1491"}
    build = re.search(r'action="(/review/[0-9_a-f]+/build)"', html).group(1)

    resp = client.post(build, data={"sku": skus}, follow_redirects=True)
    assert resp.status_code == 200
    html = resp.data.decode()
    assert "You reviewed and confirmed these columns" in html
    assert "641V19A1491" in html                # matched via the AI column
    assert "Excluded at review" not in html

    # the draft is consumed: pending dir empty, replay 404s
    assert not list((tmp_path / "data" / "pending").glob("*"))
    assert client.post(build, data={}).status_code == 404


def test_review_deselects_specific_skus(tmp_path, monkeypatch):
    """Unticked SKUs stay out of the workbook and are listed as excluded."""
    detection = [{"sheet": "Track", "header_row": 1,
                  "sku_columns": [{"column": "A", "header": "SKU",
                                   "reason": "codes"}]}]
    client, map_path, query_path = _verification_app(
        tmp_path, monkeypatch, detection)
    data, _ = _finish_job(client, _upload(client, map_path, query_path))
    html = client.get(data["next"]).data.decode()
    build = re.search(r'action="(/review/[0-9_a-f]+/build)"', html).group(1)

    # keep only the base that has stock rows
    resp = client.post(build, data={"sku": ["641V19A1491"]},
                       follow_redirects=True)
    html = resp.data.decode()
    assert "wrote <strong>1</strong> rows" in html
    assert "Excluded at review (1)" in html
    assert "M0715OUQO_M900_TU" in html          # named in the excluded list


@pytest.mark.parametrize("detection", [
    [],  # empty workbook
    # the realistic Claude answer for a no-SKU workbook: one entry per
    # sheet, each with an empty sku_columns list — must NOT pause on review
    [{"sheet": "Track", "header_row": 1, "sku_columns": []}],
])
def test_review_skipped_when_nothing_found(tmp_path, monkeypatch, detection):
    client, map_path, query_path = _verification_app(
        tmp_path, monkeypatch, detection)
    data, _ = _finish_job(client, _upload(client, map_path, query_path))
    assert data["status"] == "done", data
    # nothing to review: straight to the (empty) report
    assert "/report/" in data["next"]
    html = client.get(data["next"]).data.decode()
    assert "found no SKU columns" in html
    assert "unticked" not in html


def _marker_workbook(path, header="PickUp"):
    """One sheet with a PickUp column: marked, highlighted-only,
    unrecognized-value, and marked-but-empty rows."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Mark"
    ws["A1"] = "SKU"
    ws["B1"] = header
    ws["C1"] = "Qty"
    ws["A2"] = "644S16B7E72X5805"          # marked -> selected
    ws["B2"] = "Y"
    ws["A3"] = "641V19A1491X8300"          # highlighted only -> NOT selected
    ws["A3"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["A4"] = "652P92X3F74X8090"           # marked via variant -> selected
    ws["B4"] = "yes"
    ws["A5"] = "657H91A1234X5678"           # unrecognized value -> ignored
    ws["B5"] = "maybe?"
    ws["B6"] = "Y"                          # marked row without any SKU
    wb.save(path)


@pytest.mark.parametrize("header", ["PickUp", "pick up", "PICK-UP"])
def test_pickup_column_is_primary_selector(tmp_path, header):
    path = tmp_path / "mark.xlsx"
    _marker_workbook(path, header=header)
    scan = pipeline.scan_sell_thru_map(path)
    s = scan.sheets[0]
    assert s.uses_marker and s.marker_column == "B"
    assert s.marker_header_row == 1
    assert s.marked_rows == 3                       # Y, yes, and the empty row
    # marked rows select; the highlight alone does NOT
    assert s.marked_skus == {"644S16B7E72X5805", "652P92X3F74X8090"}
    assert s.bases == {"644S16B7E72", "652P92X3F74"}
    # the highlighted-but-unmarked SKU is surfaced, not silently dropped
    assert s.highlighted_unmarked == ["641V19A1491X8300"]
    # anomalies are recorded for the human
    assert s.marker_unrecognized == [("B5", "maybe?")]
    assert s.marker_no_sku_rows == ["B6"]
    # sheet-level union still works via MapScan
    assert set(scan.bases) == {"644S16B7E72", "652P92X3F74"}


def test_pickup_marks_ai_located_skus(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Track"
    ws["A1"] = "TS SKU"
    ws["B1"] = "PickUp"
    ws["A2"] = "M0759OWKAM912"              # non-standard format, marked
    ws["B2"] = "Y"
    ws["A3"] = "KCK554TFS_S03W"             # non-standard, unmarked
    path = tmp_path / "track.xlsx"
    wb.save(path)
    ai_columns = {"Track": {"header_row": 1, "columns": {"A": "TS SKU"}}}
    scan = pipeline.scan_sell_thru_map(path, ai_columns=ai_columns)
    s = scan.sheets[0]
    assert s.marked_skus == {"M0759OWKAM912"}
    assert s.bases == {"M0759OWKAM912"}
    # recorded in the AI trace so column memory still sees the approval
    assert [d[1] for d in s.ai_highlighted_cell_details] == ["M0759OWKAM912"]


def test_pickup_and_highlight_sheets_mix(tmp_path):
    """A PickUp sheet and a highlight-only sheet in one workbook: each
    sheet keeps its own selection rule and the bases union."""
    path = tmp_path / "mixed.xlsx"
    _marker_workbook(path)
    wb = load_workbook(path)
    ws = wb.create_sheet("Plain")
    ws["A1"] = "SKU"
    ws["A2"] = "641V19A1491X8300"
    ws["A2"].fill = PatternFill("solid", fgColor="FFFF00")
    wb.save(path)
    scan = pipeline.scan_sell_thru_map(path)
    assert set(scan.bases) == {"644S16B7E72", "652P92X3F74", "641V19A1491"}
    assert scan.bases["641V19A1491"] == ["Plain"]   # not from the Mark sheet


def test_pickup_e2e_review_badge_and_report(tmp_path):
    """Upload -> review shows the PickUp badge and warnings -> report
    documents the marker column and its anomalies."""
    app = create_app(data_dir=tmp_path / "data", open_access=True)
    client = app.test_client()
    map_path = tmp_path / "mark.xlsx"
    _marker_workbook(map_path)
    qwb = Workbook()
    qws = qwb.active
    qws.title = "query"
    qws.append(pipeline.QUERY_HEADERS)
    qws.append(["361723", "644S16B7E72 - X5805 - T36", "d", "dep", "s",
                None, None, None, 1, 10, 20, "t", "p"])
    query_path = tmp_path / "q.xlsx"
    qwb.save(query_path)

    data, _ = _finish_job(client, _upload(client, map_path, query_path))
    assert "/review/" in data["next"]
    html = client.get(data["next"]).data.decode()
    assert "badge-ok" in html and "PickUp" in html      # source badge
    skus = re.findall(r'name="sku" value="([^"]+)"', html)
    assert set(skus) == {"644S16B7E72", "652P92X3F74"}
    assert "641V19A1491" not in html                    # highlight ignored
    # the anomaly warning names all three counts
    assert "unrecognized PickUp value" in html

    build = re.search(r'action="(/review/[0-9_a-f]+/build)"', html).group(1)
    resp = client.post(build, data={"sku": skus}, follow_redirects=True)
    report = resp.data.decode()
    assert "PickUp column (explicit selection)" in report
    assert "maybe?" in report                           # unrecognized listed
    assert "B6" in report                               # no-SKU row listed
    assert "641V19A1491X8300" in report                # unmarked highlight


def test_column_memory_fingerprint_and_lookup(tmp_path):
    cells = {"A": " Capsule ", "B": "TS  SKU", "C": "Qty"}
    fp = column_memory.fingerprint_row(cells)
    # normalization: whitespace and case do not change the fingerprint
    assert fp == column_memory.fingerprint_row(
        {"A": "capsule", "B": "ts sku", "C": "QTY"})
    # a shifted layout (same headers, different letters) misses safely
    assert fp != column_memory.fingerprint_row(
        {"B": "Capsule", "C": "TS SKU", "D": "Qty"})
    # near-empty rows carry no identity
    assert column_memory.fingerprint_row({"A": "SKU"}) == ""

    store = tmp_path / "mem.json"
    entry = {"sheet": "Track", "header_row": 1, "fingerprint": fp,
             "header_cells": cells,
             "sku_columns": [{"column": "B", "header": "TS SKU",
                              "reason": "codes"}]}
    column_memory.remember(store, entry, user="albert")
    previews = [{"sheet": "Other name, same layout",
                 "rows": {"1": {"A": "CAPSULE", "B": "TS SKU", "C": "Qty"},
                          "2": {"B": "M0759OWKAM912"}}}]
    remembered, unknown = column_memory.lookup(previews, store)
    assert unknown == []
    assert remembered[0]["sku_columns"][0]["column"] == "B"
    assert remembered[0]["source"] == "memory"
    assert remembered[0]["header_row"] == 1

    # header text at the stored letter must still match
    shifted = [{"sheet": "S", "rows": {"1": {"A": "CAPSULE", "B": "Renamed",
                                             "C": "Qty"}}}]
    remembered, unknown = column_memory.lookup(shifted, store)
    assert remembered == [] and len(unknown) == 1

    column_memory.forget(store, fp)
    remembered, _ = column_memory.lookup(previews, store)
    assert remembered == []


def test_column_memory_flow(tmp_path, monkeypatch):
    """Approving a sheet's AI columns stores the mapping; the next upload of
    the same layout is served from memory without calling the AI."""
    detection = [{"sheet": "Track", "header_row": 1,
                  "sku_columns": [{"column": "A", "header": "SKU",
                                   "reason": "codes"}]}]
    calls = {"n": 0}

    def fake_locate(previews):
        calls["n"] += 1
        column_memory.attach_fingerprints(detection, previews)
        return [dict(e, source="ai") for e in detection]

    monkeypatch.setattr(sku_locator, "is_configured", lambda: True)
    monkeypatch.setattr(sku_locator, "locate_from_previews", fake_locate)
    app = create_app(data_dir=tmp_path / "data", open_access=True)
    client = app.test_client()
    map_path = tmp_path / "track.xlsx"
    _track_workbook(map_path)
    qwb = Workbook()
    qws = qwb.active
    qws.title = "query"
    qws.append(pipeline.QUERY_HEADERS)
    query_path = tmp_path / "q.xlsx"
    qwb.save(query_path)

    # run 1: AI detects, reviewer keeps the AI-found SKUs -> remembered
    data, _ = _finish_job(client, _upload(client, map_path, query_path))
    assert calls["n"] == 1
    html = client.get(data["next"]).data.decode()
    skus = re.findall(r'name="sku" value="([^"]+)"', html)
    build = re.search(r'action="(/review/[0-9_a-f]+/build)"', html).group(1)
    client.post(build, data={"sku": skus})
    assert (tmp_path / "data" / "column_memory.json").exists()

    # run 2: same layout -> columns come from memory, no second AI call
    data, _ = _finish_job(client, _upload(client, map_path, query_path,
                                          extra_query=False))
    assert calls["n"] == 1, "memory hit must not call the AI again"
    html = client.get(data["next"]).data.decode()
    assert "remembered" in html          # source badge on the review page
    assert set(re.findall(r'name="sku" value="([^"]+)"', html)) \
        == {"M0715OUQO_M900_TU", "641V19A1491"}

    # run 2 build excluding EVERY AI-found SKU -> mapping is dropped
    build = re.search(r'action="(/review/[0-9_a-f]+/build)"', html).group(1)
    client.post(build, data={})
    # run 3: memory is gone, the AI is consulted again
    data, _ = _finish_job(client, _upload(client, map_path, query_path,
                                          extra_query=False))
    assert calls["n"] == 2


def test_memory_survives_ai_failure(tmp_path, monkeypatch):
    """An API outage while detecting UNKNOWN sheets must not cost the run
    its remembered sheets — they were approved and need no API."""
    detection = [{"sheet": "Track", "header_row": 1,
                  "sku_columns": [{"column": "A", "header": "SKU",
                                   "reason": "codes"}]}]

    def fake_locate(previews):
        column_memory.attach_fingerprints(detection, previews)
        return [dict(e, source="ai") for e in detection]

    monkeypatch.setattr(sku_locator, "is_configured", lambda: True)
    monkeypatch.setattr(sku_locator, "locate_from_previews", fake_locate)
    app = create_app(data_dir=tmp_path / "data", open_access=True)
    client = app.test_client()
    map_path = tmp_path / "track.xlsx"
    _track_workbook(map_path)
    qwb = Workbook()
    qws = qwb.active
    qws.title = "query"
    qws.append(pipeline.QUERY_HEADERS)
    query_path = tmp_path / "q.xlsx"
    qwb.save(query_path)

    # run 1: detect + approve -> Track's mapping is remembered
    data, _ = _finish_job(client, _upload(client, map_path, query_path))
    html = client.get(data["next"]).data.decode()
    skus = re.findall(r'name="sku" value="([^"]+)"', html)
    build = re.search(r'action="(/review/[0-9_a-f]+/build)"', html).group(1)
    client.post(build, data={"sku": skus})

    # run 2: same Track layout plus an unknown sheet, with the API down
    def boom(previews):
        raise sku_locator.SkuLocatorError("simulated outage")
    monkeypatch.setattr(sku_locator, "locate_from_previews", boom)
    map2 = tmp_path / "track2.xlsx"
    _track_workbook(map2)
    wb = load_workbook(map2)
    ws = wb.create_sheet("Mystery")
    ws["A1"] = "Something"
    ws["B1"] = "Else"
    ws["A2"] = "value-1"
    wb.save(map2)

    data, _ = _finish_job(client, _upload(client, map2, query_path,
                                          extra_query=False))
    assert data["status"] == "done", data
    assert "/review/" in data["next"], "remembered sheet must reach review"
    html = client.get(data["next"]).data.decode()
    assert "remembered" in html                  # memory badge still shown
    assert "simulated outage" in html            # failure still reported
    # the remembered column's SKUs were extracted despite the outage
    assert set(re.findall(r'name="sku" value="([^"]+)"', html)) \
        == {"M0715OUQO_M900_TU", "641V19A1491"}


def test_named_users_auth_and_identity(tmp_path, monkeypatch):
    """Per-person session logins; runs record who uploaded and who built."""
    detection = [{"sheet": "Track", "header_row": 1,
                  "sku_columns": [{"column": "A", "header": "SKU",
                                   "reason": "codes"}]}]
    monkeypatch.setattr(sku_locator, "is_configured", lambda: True)
    monkeypatch.setattr(sku_locator, "locate_from_previews",
                        lambda previews: detection)
    app = create_app(data_dir=tmp_path / "data", setup_code="root-code",
                     secret="s" * 64, cookie_secure=False)
    albert = app.test_client()
    # bootstrap the admin account (the display name feeds the report)
    r = albert.post("/setup", data={"code": "root-code", "username": "albert",
                                    "password": "pw1secret",
                                    "display": "Albert"})
    assert r.status_code == 303
    assert albert.get("/").status_code == 200
    # anonymous requests bounce to login now that a user exists
    anon = app.test_client()
    r = anon.get("/")
    assert r.status_code == 303 and r.headers["Location"].endswith("/login")
    # albert adds vivian with an initial password (manual path)
    r = albert.post("/team/add", data={"username": "vivian",
                                       "password": "pw2secret"})
    assert "msg_user_added" in r.headers["Location"]
    vivian = app.test_client()
    assert vivian.post("/login", data={"username": "vivian",
                                       "password": "wrong-pass"}
                       ).status_code == 401
    # correct login, case-insensitive username
    r = vivian.post("/login", data={"username": "VIVIAN",
                                    "password": "pw2secret"})
    assert r.status_code == 303
    assert vivian.get("/").status_code == 200

    map_path = tmp_path / "track.xlsx"
    _track_workbook(map_path)
    qwb = Workbook()
    qws = qwb.active
    qws.title = "query"
    qws.append(pipeline.QUERY_HEADERS)
    query_path = tmp_path / "q.xlsx"
    qwb.save(query_path)

    # albert uploads; vivian reviews and builds
    with map_path.open("rb") as m, query_path.open("rb") as q:
        resp = albert.post(
            "/process",
            data={"map_file": (m, "t.xlsx"), "query_file": (q, "q.xlsx")},
            content_type="multipart/form-data")
    assert resp.status_code == 302
    job = resp.headers["Location"].rsplit("/", 1)[-1]
    while True:
        d = albert.get(f"/progress/{job}/status").get_json()
        if d["status"] != "running":
            break
        time.sleep(0.2)
    html = vivian.get(d["next"]).data.decode()
    skus = re.findall(r'name="sku" value="([^"]+)"', html)
    build = re.search(r'action="(/review/[0-9_a-f]+/build)"', html).group(1)
    resp = vivian.post(build, data={"sku": skus}, follow_redirects=True)
    html = resp.data.decode()
    assert "Uploaded by Albert" in html          # display name, not username
    assert "Reviewed &amp; built by vivian" in html
    # the recent-runs list names the builder too
    index = albert.get("/").data.decode()
    assert "vivian" in index


def test_review_build_uses_rows_matched_at_scan_time(tmp_path, monkeypatch):
    """Replacing the stored query while a review page sits open must not
    change the data of the run being reviewed — matching happened at scan
    time and the draft carries the rows."""
    detection = [{"sheet": "Track", "header_row": 1,
                  "sku_columns": [{"column": "A", "header": "SKU",
                                   "reason": "codes"}]}]
    client, map_path, _ = _verification_app(tmp_path, monkeypatch, detection)

    def make_query(path, cost, retail):
        wb = Workbook()
        ws = wb.active
        ws.title = "query"
        ws.append(pipeline.QUERY_HEADERS)
        ws.append(["3617000000001", "641V19A1491 - X8300 - T36", "d", "dep",
                   "s", None, None, None, 1, cost, retail, "t", "p"])
        wb.save(path)

    query_a = tmp_path / "qa.xlsx"
    make_query(query_a, 123.45, 246.9)
    query_b = tmp_path / "qb.xlsx"
    make_query(query_b, 987.65, 1975.3)

    # user A scans with query A and pauses on the review page
    data, _ = _finish_job(client, _upload(client, map_path, query_a))
    review_url = data["next"]

    # meanwhile user B replaces the stored query with query B
    _finish_job(client, _upload(client, map_path, query_b))

    # A builds: the run must still carry query A's numbers
    html = client.get(review_url).data.decode()
    skus = re.findall(r'name="sku" value="([^"]+)"', html)
    build = re.search(r'action="(/review/[0-9_a-f]+/build)"', html).group(1)
    resp = client.post(build, data={"sku": skus}, follow_redirects=True)
    html = resp.data.decode()
    assert "123.45" in html
    assert "987.65" not in html


@pytest.mark.skipif(
    not sku_locator.is_configured()
    or not (SAMPLES / "delivery_track.xlsx").exists(),
    reason="needs ANTHROPIC_API_KEY and samples/delivery_track.xlsx",
)
def test_locator_live_on_delivery_track():
    detection = sku_locator.locate(SAMPLES / "delivery_track.xlsx")
    by_sheet = {d["sheet"]: [c["column"] for c in d["sku_columns"]]
                for d in detection}
    assert "E" in by_sheet.get("All Cat", [])   # 'TS SKU' column
    assert "A" in by_sheet.get("Sheet2", [])    # 'SKU' column


# --- 3. query matching ------------------------------------------------------

@needs_samples
def test_every_fw26_base_matches_query(map_scan, fw26_matched):
    fw26 = next(s for s in map_scan.sheets if s.name == "FW26")
    assert set(fw26_matched) == fw26.bases  # every base has >= 1 row
    assert all(len(rows) >= 1 for rows in fw26_matched.values())
    total = sum(len(rows) for rows in fw26_matched.values())
    # "~283" in the spec; the exact figure shifts by a row or two between
    # query export versions, so assert a tight band, not one copy's value.
    assert 275 <= total <= 295


def test_match_is_exact_first_segment(tmp_path):
    """'644S16B7E7' must not prefix-match '644S16B7E72 - ...'."""
    wb = Workbook()
    ws = wb.active
    ws.title = "query"
    ws.append(pipeline.QUERY_HEADERS)
    ws.append(["1", "644S16B7E72 - X4804 - T36", "d", "dep", "s",
               None, None, None, 1, 10, 20, "t", "p"])
    path = tmp_path / "q.xlsx"
    wb.save(path)
    assert pipeline.match_query(path, {"644S16B7E7"}) == {}
    assert list(pipeline.match_query(path, {"644S16B7E72"})) == ["644S16B7E72"]


# --- 4. output workbook -----------------------------------------------------

@needs_samples
def test_output_workbook_structure(tmp_path, fw26_matched):
    out = tmp_path / "out.xlsx"
    rows_written = pipeline.build_output(
        TEMPLATE, out, fw26_matched,
        unmatched_bases=["FAKEBASE001"],
        base_sheets={"FAKEBASE001": ["LOOK TOTAL"]},
    )
    # every matched query row must land in the workbook, no more, no fewer
    assert rows_written == sum(len(rows) for rows in fw26_matched.values())

    wb = load_workbook(out)
    ws = wb["Sheet1"]
    twb = load_workbook(TEMPLATE)
    assert [c.value for c in ws[1]] == [c.value for c in twb["Sheet1"][1]]
    twb.close()

    last = rows_written + 1
    for n in (2, 3, last):
        assert ws.cell(row=n, column=8).value == \
            f'=IF(AND(F{n}<>"", G{n}<>""), F{n}*G{n}, "")'
        assert ws.cell(row=n, column=10).value == \
            f'=IF(AND(F{n}<>"", I{n}<>""), F{n}*I{n}, "")'
        for col in (11, 12, 13):  # K, L, M stay blank
            assert ws.cell(row=n, column=col).value is None
        barcode = ws.cell(row=n, column=1)
        assert isinstance(barcode.value, str)
        assert barcode.number_format == "@"

    # sorted by base, then color, then size
    skus = [ws.cell(row=n, column=2).value for n in range(2, last + 1)]
    keys = [tuple(s.split(" - ")) for s in skus]
    assert keys == sorted(keys)

    um = wb["Unmatched"]
    assert um["A1"].value == "Base SKU"
    assert um["A2"].value == "FAKEBASE001"
    assert um["B2"].value == "LOOK TOTAL"
    wb.close()


# --- 5. Flask end-to-end ----------------------------------------------------

@needs_samples
def test_end_to_end_upload_and_download(tmp_path):
    app = create_app(data_dir=tmp_path, open_access=True)
    client = app.test_client()

    page = client.get("/")
    assert page.status_code == 200
    assert b"No stock query export stored yet" in page.data

    with MAP.open("rb") as m, QUERY.open("rb") as q:
        resp = client.post(
            "/process",
            data={"map_file": (m, "map.xlsx"), "query_file": (q, "query.xlsx")},
            content_type="multipart/form-data",
        )
    # the upload hands off to a background job with a polled progress bar
    data, statuses = _finish_job(client, resp)
    assert data["status"] == "done", data
    # on a real-sized workbook the bar reports true intermediate progress
    assert any(0 < (s.get("percent") or 0) < 100 for s in statuses)

    # the scan lands on the review checkpoint; keep every SKU and build
    assert "/review/" in data["next"]
    review_html = client.get(data["next"]).data.decode()
    skus = re.findall(r'name="sku" value="([^"]+)"', review_html)
    assert len(skus) == 41
    build_url = re.search(r'action="(/review/[0-9_a-f]+/build)"',
                          review_html).group(1)
    resp = client.post(build_url, data={"sku": skus})
    assert resp.status_code == 302
    resp = client.get(resp.headers["Location"])
    assert resp.status_code == 200
    html = resp.data.decode()
    match = re.search(r"/download/(\d{8}_\d{6}_[0-9a-f]{6})", html)
    assert match, "report page must contain a download link"
    rows = re.search(r"wrote <strong>(\d+)</strong> rows", html)
    assert rows, "report page must state the number of rows written"
    rows_written = int(rows.group(1))
    assert 275 <= rows_written <= 295  # "~283", varies with query version

    # the report must trace internal steps: highlighted cells found, base
    # extraction, and the query data pulled per base
    assert "Processing steps" in html
    # without an API key configured, the AI section says it is off
    assert "AI column detection" in html
    cell_trace = re.search(
        r"<td>([A-Z]{1,3}\d+)</td><td>([0-9A-Z]{8,})</td><td>([0-9A-Z]{8,})</td>",
        html.replace("\n", ""))
    assert cell_trace, "step 1 must list highlighted cells with SKU and base"
    detail = re.search(r"<summary>([0-9A-Z]{8,}) — (\d+) query", html)
    assert detail, "step 3 must show per-base query detail"
    assert detail.group(1) in html  # base appears in the matched table too

    dl = client.get(f"/download/{match.group(1)}")
    assert dl.status_code == 200
    assert "ProductsList_" in dl.headers["Content-Disposition"]
    wb = load_workbook(io.BytesIO(dl.data))
    assert "Sheet1" in wb.sheetnames and "Unmatched" in wb.sheetnames
    # the downloaded workbook must contain exactly the reported rows
    assert wb["Sheet1"].max_row == rows_written + 1
    wb.close()

    # the query is now stored: a map-only run must also succeed, and get
    # its own run_id even within the same second (concurrent-run safety)
    page = client.get("/")
    assert b"Stored stock query" in page.data
    with MAP.open("rb") as m:
        resp = client.post("/process", data={"map_file": (m, "map.xlsx")},
                           content_type="multipart/form-data")
    second, _ = _finish_job(client, resp)
    assert second["status"] == "done"
    # a second upload gets its own review draft, distinct from the first
    assert "/review/" in second["next"]
    assert second["next"] != data["next"]

    # the report is persistent (refresh-safe) and language-switchable:
    # the same run re-renders in Chinese with the same download link
    zh = client.get(f"/report/{match.group(1)}?lang=zh")
    assert zh.status_code == 200
    zh_html = zh.data.decode()
    assert "处理步骤" in zh_html            # "Processing steps"
    assert "标黄基础款号" in zh_html         # cohesive summary wording
    assert f"/download/{match.group(1)}" in zh_html
    # and back to English via the toggle
    en = client.get(f"/report/{match.group(1)}?lang=en")
    assert "Processing steps" in en.data.decode()


def test_query_upload_with_wrong_headers_is_rejected(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "query"
    ws.append(["Totally", "Wrong", "Headers"])
    bad = tmp_path / "bad_query.xlsx"
    wb.save(bad)

    ok, message, _ = pipeline.validate_query_file(bad)
    assert not ok and "header" in message

    app = create_app(data_dir=tmp_path / "data", open_access=True)
    client = app.test_client()
    with bad.open("rb") as b, bad.open("rb") as m:
        resp = client.post(
            "/process",
            data={"map_file": (m, "map.xlsx"), "query_file": (b, "q.xlsx")},
            content_type="multipart/form-data",
        )
    assert resp.status_code == 400
    assert b"rejected" in resp.data


def test_malformed_map_upload_is_a_friendly_error(tmp_path):
    app = create_app(data_dir=tmp_path, open_access=True)
    client = app.test_client()
    # store a minimal valid query first
    wb = Workbook()
    ws = wb.active
    ws.title = "query"
    ws.append(pipeline.QUERY_HEADERS)
    good_q = tmp_path / "q.xlsx"
    wb.save(good_q)
    with good_q.open("rb") as q:
        resp = client.post(
            "/process",
            data={"map_file": (io.BytesIO(b"this is not a zip"), "map.xlsx"),
                  "query_file": (q, "q.xlsx")},
            content_type="multipart/form-data",
        )
    assert resp.status_code == 400
    assert b"could not be processed" in resp.data


def test_session_required_when_setup_code_set(tmp_path):
    app = create_app(data_dir=tmp_path, setup_code="root-code",
                     secret="s" * 64, cookie_secure=False)
    client = app.test_client()
    # zero users: anonymous requests bounce to first-run setup
    r = client.get("/")
    assert r.status_code == 303 and r.headers["Location"].endswith("/setup")
    r = client.post("/setup", data={"code": "root-code", "username": "boss",
                                    "password": "longenough"})
    assert r.status_code == 303
    assert client.get("/").status_code == 200
    # the setup code is NOT a login password
    fresh = app.test_client()
    assert fresh.post("/login", data={"username": "boss",
                                      "password": "root-code"}
                      ).status_code == 401
    assert fresh.post("/login", data={"username": "boss",
                                      "password": "longenough"}
                      ).status_code == 303


def test_language_toggle_and_cookie(tmp_path):
    app = create_app(data_dir=tmp_path, open_access=True)
    client = app.test_client()

    # default is English, with a top-left switch offering Chinese
    page = client.get("/")
    assert "No stock query export stored yet" in page.data.decode()
    assert 'class="lang"' in page.data.decode()
    assert "中文" in page.data.decode()

    # ?lang=zh renders Chinese and persists the choice in a cookie
    page = client.get("/?lang=zh")
    html = page.data.decode()
    assert "还没有保存库存查询表" in html
    assert "开始处理" in html          # Process button
    assert "English" in html           # switch now offers English
    assert "lang=zh" in page.headers.get("Set-Cookie", "")

    # subsequent plain requests stay in Chinese via the cookie
    page = client.get("/")
    assert "还没有保存库存查询表" in page.data.decode()

    # switching back works
    page = client.get("/?lang=en")
    assert "No stock query export stored yet" in page.data.decode()

    # error pages are translated too
    resp = client.post("/process?lang=zh", data={},
                       content_type="multipart/form-data")
    assert resp.status_code == 400
    assert "请先选择 Sell-Thru Map" in resp.data.decode()


def test_help_page_in_both_languages(tmp_path):
    app = create_app(data_dir=tmp_path, open_access=True)
    client = app.test_client()

    # linked from the front page
    assert b"/help" in client.get("/").data

    en = client.get("/help").data.decode()
    assert en.count("<h2>") >= 5
    assert "What this tool does" in en
    assert "Your first run, step by step" in en
    assert "filter" in en  # covers the hidden-rows gotcha

    zh = client.get("/help?lang=zh").data.decode()
    assert "这个工具是做什么的" in zh
    assert "第一次使用" in zh
    assert "筛选" in zh

    # help is behind auth like everything else
    locked = create_app(data_dir=tmp_path / "locked", setup_code="pw",
                        cookie_secure=False)
    assert locked.test_client().get("/help").status_code == 303


def test_corrupt_xml_map_gets_friendly_error(tmp_path):
    """A valid zip with truncated sheet XML raises ElementTree.ParseError —
    it must land on the friendly error page, not a bare 500."""
    import zipfile as zf

    src = Path(__file__).parent / "assets" / "ProductsListTemplate.xlsx"
    corrupt = io.BytesIO()
    with zf.ZipFile(src) as zin, zf.ZipFile(corrupt, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/"):
                data = data[: len(data) // 2]
            zout.writestr(item, data)
    corrupt.seek(0)

    app = create_app(data_dir=tmp_path, open_access=True)
    client = app.test_client()
    wb = Workbook()
    ws = wb.active
    ws.title = "query"
    ws.append(pipeline.QUERY_HEADERS)
    good_q = tmp_path / "q.xlsx"
    wb.save(good_q)
    with good_q.open("rb") as q:
        resp = client.post(
            "/process",
            data={"map_file": (corrupt, "map.xlsx"), "query_file": (q, "q.xlsx")},
            content_type="multipart/form-data",
        )
    # it passes the quick zip check, so the failure surfaces on the
    # progress page as a job error with the parse detail — never a dead bar
    data, _ = _finish_job(client, resp)
    assert data["status"] == "error"
    assert data["detail"]


def test_xhr_submission_gets_json_contract(tmp_path):
    """The upload-progress path posts with XHR and expects JSON back:
    {'next': url} on success, {'error': message} on failure."""
    app = create_app(data_dir=tmp_path, open_access=True)
    client = app.test_client()
    xhr = {"X-Requested-With": "XMLHttpRequest"}

    # error path: no map file selected
    resp = client.post("/process", data={}, headers=xhr,
                       content_type="multipart/form-data")
    assert resp.status_code == 400
    assert "error" in resp.get_json()

    # success path: tiny valid map + query -> {'next': /progress/<job>}
    wb = Workbook()
    ws = wb.active
    ws.title = "query"
    ws.append(pipeline.QUERY_HEADERS)
    q = tmp_path / "q.xlsx"
    wb.save(q)
    m = tmp_path / "m.xlsx"
    _track_workbook(m)
    with m.open("rb") as mf, q.open("rb") as qf:
        resp = client.post(
            "/process",
            data={"map_file": (mf, "m.xlsx"), "query_file": (qf, "q.xlsx")},
            headers=xhr, content_type="multipart/form-data")
    assert resp.status_code == 200
    nxt = resp.get_json()["next"]
    assert "/progress/" in nxt
    data, _ = _wait_status(client, nxt.rsplit("/", 1)[-1])
    assert data["status"] == "done"


def test_indexed_color_fills_are_flagged_not_dropped(tmp_path):
    """Legacy indexed-palette fills are outside the highlight definition but
    must surface in 'other fills detected' instead of vanishing silently."""
    wb = Workbook()
    ws = wb.active
    c = ws.cell(row=1, column=1, value="644T000A99X5800")
    c.fill = PatternFill(fill_type="solid", start_color=Color(indexed=6))
    path = tmp_path / "indexed.xlsx"
    wb.save(path)
    scan = pipeline.scan_sell_thru_map(path)
    sheet = scan.sheets[0]
    assert sheet.highlighted_cells == 0
    assert len(sheet.other_fills) == 1
    assert "indexed" in sheet.other_fills[0][2]
